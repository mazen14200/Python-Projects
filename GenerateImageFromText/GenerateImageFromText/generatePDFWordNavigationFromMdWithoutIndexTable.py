import re
import os
import sys
import html
import tempfile
from pathlib import Path
from typing import Optional

import fitz  # PyMuPDF
import win32com.client

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
)
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


# ============================================================
# Configuration
# ============================================================

BLUE_ENGLISH = '3662AE'
BROWN_ENGLISH = 'AE6A36'
BODY_COLOR = '2C3E50'

HEADING_COLORS = {
    1: '1A5276',
    2: '21618C',
    3: '2E86C1',
    4: '3498DB',
}

TABLE_HEADER_COLOR = '2980B9'
TABLE_BORDER_COLOR = 'D5DBDB'
TOC_HEADER_COLOR = '2980B9'


# ============================================================
# Regular Expressions
# ============================================================

ARABIC_RE = re.compile(
    r'[\u0600-\u06FF'
    r'\u0750-\u077F'
    r'\u08A0-\u08FF'
    r'\uFB50-\uFDFF'
    r'\uFE70-\uFEFF]'
)

LATIN_RE = re.compile(r'[A-Za-z]')

URL_RE = re.compile(
    r'https?://\S+$',
    re.IGNORECASE
)

MARKDOWN_LINK_RE = re.compile(
    r'^\[([^\]]+)\]\(([^)]+)\)$'
)

INLINE_LINK_RE = re.compile(
    r'\[([^\]]+)\]\((https?://[^)]+)\)',
    re.IGNORECASE
)

QUESTION_MARKER_RE = re.compile(
    r'\b[qQ]\d+\s*:'
)


# ============================================================
# Text / Font Helpers
# ============================================================

def add_colored_text_run(
    paragraph,
    text,
    *,
    size=11,
    bold=False,
    italic=False,
    default_color=BODY_COLOR
):
    """
    Add text while forcing Q<number>: markers to blue.
    """

    if not text:
        return

    pos = 0

    for match in QUESTION_MARKER_RE.finditer(text):

        if match.start() > pos:
            chunk = text[pos:match.start()]

            if chunk:
                run = paragraph.add_run(chunk)

                set_run_font(
                    run,
                    size=size,
                    bold=bold,
                    italic=italic,
                    color=default_color
                )

        run = paragraph.add_run(match.group(0))

        set_run_font(
            run,
            size=size,
            bold=bold,
            italic=italic,
            color=BLUE_ENGLISH
        )

        pos = match.end()

    if pos < len(text):

        chunk = text[pos:]

        run = paragraph.add_run(chunk)

        set_run_font(
            run,
            size=size,
            bold=bold,
            italic=italic,
            color=default_color
        )


def read_markdown_file(file_path):

    encodings = [
        'utf-8-sig',
        'utf-8',
        'cp1256',
        'windows-1256',
        'utf-16'
    ]

    for enc in encodings:

        try:

            with open(
                file_path,
                'r',
                encoding=enc
            ) as f:

                return f.read().lstrip('\ufeff')

        except UnicodeDecodeError:

            continue

    with open(
        file_path,
        'r',
        encoding='utf-8',
        errors='replace'
    ) as f:

        return f.read().lstrip('\ufeff')


def html_escape(text):
    return html.escape(
        str(text),
        quote=False
    )


def normalize_url(url: str) -> str:
    return html.unescape(url).strip()


def is_english_only_text(text):

    if not text or not text.strip():
        return False

    if ARABIC_RE.search(text):
        return False

    return bool(
        LATIN_RE.search(text)
    )


def english_color(text):

    return (
        BLUE_ENGLISH
        if '?' in text
        else BROWN_ENGLISH
    )


# ============================================================
# Word XML Helpers
# ============================================================

def add_bookmark(
    paragraph,
    name: str,
    bookmark_id: int
):

    start = OxmlElement(
        'w:bookmarkStart'
    )

    start.set(
        qn('w:id'),
        str(bookmark_id)
    )

    start.set(
        qn('w:name'),
        name
    )

    end = OxmlElement(
        'w:bookmarkEnd'
    )

    end.set(
        qn('w:id'),
        str(bookmark_id)
    )

    paragraph._p.insert(
        0,
        start
    )

    paragraph._p.append(
        end
    )


def add_internal_hyperlink(
    paragraph,
    text: str,
    anchor: str,
    *,
    bold=False,
    color='2980B9',
    underline=False
):

    hyperlink = OxmlElement(
        'w:hyperlink'
    )

    hyperlink.set(
        qn('w:anchor'),
        anchor
    )

    run = OxmlElement(
        'w:r'
    )

    rPr = OxmlElement(
        'w:rPr'
    )

    if bold:

        b = OxmlElement(
            'w:b'
        )

        rPr.append(b)

    color_el = OxmlElement(
        'w:color'
    )

    color_el.set(
        qn('w:val'),
        color
    )

    rPr.append(
        color_el
    )

    if underline:

        u = OxmlElement(
            'w:u'
        )

        u.set(
            qn('w:val'),
            'single'
        )

        rPr.append(u)

    rFonts = OxmlElement(
        'w:rFonts'
    )

    rFonts.set(
        qn('w:ascii'),
        'Segoe UI'
    )

    rFonts.set(
        qn('w:hAnsi'),
        'Segoe UI'
    )

    rFonts.set(
        qn('w:cs'),
        'Arial'
    )

    rPr.append(
        rFonts
    )

    run.append(
        rPr
    )

    t = OxmlElement(
        'w:t'
    )

    t.text = text

    run.append(
        t
    )

    hyperlink.append(
        run
    )

    paragraph._p.append(
        hyperlink
    )

    return hyperlink


def add_external_hyperlink(
    paragraph,
    text: str,
    url: str,
    *,
    color=None,
    underline=False,
    bold=False
):

    url = normalize_url(url)

    rid = paragraph.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True
    )

    hyperlink = OxmlElement(
        'w:hyperlink'
    )

    hyperlink.set(
        qn('r:id'),
        rid
    )

    run = OxmlElement(
        'w:r'
    )

    rPr = OxmlElement(
        'w:rPr'
    )

    if bold:

        b = OxmlElement(
            'w:b'
        )

        rPr.append(b)

    if color:

        c = OxmlElement(
            'w:color'
        )

        c.set(
            qn('w:val'),
            color
        )

        rPr.append(c)

    if underline:

        u = OxmlElement(
            'w:u'
        )

        u.set(
            qn('w:val'),
            'single'
        )

        rPr.append(u)

    rFonts = OxmlElement(
        'w:rFonts'
    )

    rFonts.set(
        qn('w:ascii'),
        'Segoe UI'
    )

    rFonts.set(
        qn('w:hAnsi'),
        'Segoe UI'
    )

    rFonts.set(
        qn('w:cs'),
        'Arial'
    )

    rPr.append(
        rFonts
    )

    run.append(
        rPr
    )

    t = OxmlElement(
        'w:t'
    )

    t.text = text

    run.append(
        t
    )

    hyperlink.append(
        run
    )

    paragraph._p.append(
        hyperlink
    )

    return hyperlink


def add_field(
    paragraph,
    instruction: str,
    text_fallback: str = ''
):

    run = paragraph.add_run()

    fld_begin = OxmlElement(
        'w:fldChar'
    )

    fld_begin.set(
        qn('w:fldCharType'),
        'begin'
    )

    instr = OxmlElement(
        'w:instrText'
    )

    instr.set(
        qn('xml:space'),
        'preserve'
    )

    instr.text = instruction

    fld_sep = OxmlElement(
        'w:fldChar'
    )

    fld_sep.set(
        qn('w:fldCharType'),
        'separate'
    )

    text = OxmlElement(
        'w:t'
    )

    text.text = text_fallback

    fld_end = OxmlElement(
        'w:fldChar'
    )

    fld_end.set(
        qn('w:fldCharType'),
        'end'
    )

    run._r.append(
        fld_begin
    )

    run._r.append(
        instr
    )

    run._r.append(
        fld_sep
    )

    run._r.append(
        text
    )

    run._r.append(
        fld_end
    )

    return run


def set_run_font(
    run,
    font_name='Segoe UI',
    size=None,
    bold=None,
    italic=None,
    color=None
):

    run.font.name = font_name

    rpr = run._element.get_or_add_rPr()

    rfonts = rpr.rFonts

    if rfonts is None:

        rfonts = OxmlElement(
            'w:rFonts'
        )

        rpr.append(
            rfonts
        )

    rfonts.set(
        qn('w:ascii'),
        font_name
    )

    rfonts.set(
        qn('w:hAnsi'),
        font_name
    )

    rfonts.set(
        qn('w:cs'),
        'Arial'
    )

    if size is not None:
        run.font.size = Pt(size)

    if bold is not None:
        run.bold = bold

    if italic is not None:
        run.italic = italic

    if color:

        run.font.color.rgb = (
            RGBColor.from_string(color)
        )

    return run


# ============================================================
# Table Helpers
# ============================================================

def set_cell_shading(
    cell,
    fill
):

    tcPr = cell._tc.get_or_add_tcPr()

    shd = tcPr.find(
        qn('w:shd')
    )

    if shd is None:

        shd = OxmlElement(
            'w:shd'
        )

        tcPr.append(shd)

    shd.set(
        qn('w:fill'),
        fill
    )


def set_cell_border(
    cell,
    color=TABLE_BORDER_COLOR,
    sz='6'
):

    tcPr = cell._tc.get_or_add_tcPr()

    borders = tcPr.first_child_found_in(
        'w:tcBorders'
    )

    if borders is None:

        borders = OxmlElement(
            'w:tcBorders'
        )

        tcPr.append(
            borders
        )

    for edge in (
        'top',
        'left',
        'bottom',
        'right',
        'insideH',
        'insideV'
    ):

        tag = 'w:' + edge

        el = borders.find(
            qn(tag)
        )

        if el is None:

            el = OxmlElement(tag)

            borders.append(el)

        el.set(
            qn('w:val'),
            'single'
        )

        el.set(
            qn('w:sz'),
            sz
        )

        el.set(
            qn('w:space'),
            '0'
        )

        el.set(
            qn('w:color'),
            color
        )


def set_repeat_table_header(row):

    trPr = row._tr.get_or_add_trPr()

    tblHeader = OxmlElement(
        'w:tblHeader'
    )

    tblHeader.set(
        qn('w:val'),
        'true'
    )

    trPr.append(
        tblHeader
    )


# ============================================================
# Document Settings
# ============================================================

def set_update_fields_on_open(doc):

    settings = doc.settings.element

    update = settings.find(
        qn('w:updateFields')
    )

    if update is None:

        update = OxmlElement(
            'w:updateFields'
        )

        settings.append(
            update
        )

    update.set(
        qn('w:val'),
        'true'
    )


def set_section_page_number_restart(
    section,
    start=1
):

    sectPr = section._sectPr

    pgNumType = sectPr.find(
        qn('w:pgNumType')
    )

    if pgNumType is None:

        pgNumType = OxmlElement(
            'w:pgNumType'
        )

        sectPr.append(
            pgNumType
        )

    pgNumType.set(
        qn('w:start'),
        str(start)
    )


def add_footer_page_number(section):

    footer = section.footer

    p = footer.paragraphs[0]

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    p.clear()

    run = add_field(
        p,
        'PAGE',
        '1'
    )

    set_run_font(
        run,
        size=18,
        bold=True,
        color='000000'
    )


# ============================================================
# Styles
# ============================================================

def configure_styles(doc: Document):

    styles = doc.styles

    normal = styles['Normal']

    normal.font.name = 'Segoe UI'
    normal.font.size = Pt(11)
    normal.font.color.rgb = (
        RGBColor.from_string(
            BODY_COLOR
        )
    )

    normal._element.rPr.rFonts.set(
        qn('w:ascii'),
        'Segoe UI'
    )

    normal._element.rPr.rFonts.set(
        qn('w:hAnsi'),
        'Segoe UI'
    )

    normal._element.rPr.rFonts.set(
        qn('w:cs'),
        'Arial'
    )

    for level in range(1, 5):

        style = styles[
            f'Heading {level}'
        ]

        style.font.name = 'Segoe UI'

        style.font.size = Pt(
            {
                1: 18,
                2: 15,
                3: 13,
                4: 11
            }[level]
        )

        style.font.bold = True

        style.font.color.rgb = (
            RGBColor.from_string(
                HEADING_COLORS[level]
            )
        )

        style._element.rPr.rFonts.set(
            qn('w:ascii'),
            'Segoe UI'
        )

        style._element.rPr.rFonts.set(
            qn('w:hAnsi'),
            'Segoe UI'
        )

        style._element.rPr.rFonts.set(
            qn('w:cs'),
            'Arial'
        )

    if 'Code Block' not in styles:

        style = styles.add_style(
            'Code Block',
            WD_STYLE_TYPE.PARAGRAPH
        )

    else:

        style = styles[
            'Code Block'
        ]

    style.font.name = 'Consolas'
    style.font.size = Pt(9)
    style.font.color.rgb = (
        RGBColor.from_string(
            BODY_COLOR
        )
    )

    style._element.rPr.rFonts.set(
        qn('w:ascii'),
        'Consolas'
    )

    style._element.rPr.rFonts.set(
        qn('w:hAnsi'),
        'Consolas'
    )

    style._element.rPr.rFonts.set(
        qn('w:cs'),
        'Arial'
    )


# ============================================================
# Mixed Text
# ============================================================

def add_mixed_text_runs(
    paragraph,
    text: str,
    *,
    force_color=None
):
    """
    Add text while coloring English chunks.

    English:
        ? -> blue
        normal -> brown

    Q<number>: markers are always blue.
    Arabic remains body color.
    """

    text = str(text)

    chunks = re.split(
        r'([\u0600-\u06FF'
        r'\u0750-\u077F'
        r'\u08A0-\u08FF'
        r'\uFB50-\uFDFF'
        r'\uFE70-\uFEFF]+)',
        text
    )

    for chunk in chunks:

        if not chunk:
            continue

        if (
            LATIN_RE.search(chunk)
            and not ARABIC_RE.search(chunk)
        ):

            color = (
                force_color
                or english_color(chunk)
            )

            pos = 0

            for match in QUESTION_MARKER_RE.finditer(
                chunk
            ):

                if match.start() > pos:

                    before = chunk[
                        pos:match.start()
                    ]

                    run = paragraph.add_run(
                        before
                    )

                    set_run_font(
                        run,
                        color=color,
                        bold=True,
                        size=11
                    )

                run = paragraph.add_run(
                    match.group(0)
                )

                set_run_font(
                    run,
                    color=BLUE_ENGLISH,
                    bold=True,
                    size=11
                )

                pos = match.end()

            if pos < len(chunk):

                run = paragraph.add_run(
                    chunk[pos:]
                )

                set_run_font(
                    run,
                    color=color,
                    bold=True,
                    size=11
                )

        else:

            pos = 0

            for match in QUESTION_MARKER_RE.finditer(
                chunk
            ):

                if match.start() > pos:

                    run = paragraph.add_run(
                        chunk[pos:match.start()]
                    )

                    set_run_font(
                        run,
                        color=BODY_COLOR,
                        size=11
                    )

                run = paragraph.add_run(
                    match.group(0)
                )

                set_run_font(
                    run,
                    color=BLUE_ENGLISH,
                    bold=False,
                    size=11
                )

                pos = match.end()

            if pos < len(chunk):

                run = paragraph.add_run(
                    chunk[pos:]
                )

                set_run_font(
                    run,
                    color=BODY_COLOR,
                    size=11
                )


# ============================================================
# Markdown Inline Parsing
# ============================================================

def parse_inline_segments(text: str):
    """
    Return:
        (kind, value, optional_url)
    """

    text = text.replace(
        '\u00a0',
        ' '
    )

    segments = []

    pos = 0

    combined = re.compile(
        r'\[([^\]]+)\]\((https?://[^)]+)\)'
        r'|`([^`]+)`'
        r'|\*\*([^*]+)\*\*'
        r'|\*([^*]+)\*',
        re.IGNORECASE
    )

    for m in combined.finditer(text):

        if m.start() > pos:

            segments.append(
                (
                    'text',
                    text[pos:m.start()],
                    None
                )
            )

        if m.group(1) is not None:

            segments.append(
                (
                    'link',
                    m.group(1),
                    normalize_url(
                        m.group(2)
                    )
                )
            )

        elif m.group(3) is not None:

            segments.append(
                (
                    'code',
                    m.group(3),
                    None
                )
            )

        elif m.group(4) is not None:

            segments.append(
                (
                    'bold',
                    m.group(4),
                    None
                )
            )

        else:

            segments.append(
                (
                    'italic',
                    m.group(5),
                    None
                )
            )

        pos = m.end()

    if pos < len(text):

        segments.append(
            (
                'text',
                text[pos:],
                None
            )
        )

    return segments


# ============================================================
# Rich Paragraph
# ============================================================

def add_rich_paragraph(
    doc,
    text: str,
    *,
    style=None,
    alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    hyperlink_url: Optional[str] = None
):

    p = (
        doc.add_paragraph(style=style)
        if style
        else doc.add_paragraph()
    )

    p.alignment = alignment

    p.paragraph_format.space_after = Pt(7)

    if hyperlink_url:

        label = text.strip()

        add_external_hyperlink(
            p,
            label,
            hyperlink_url,
            color=english_color(label),
            bold=is_english_only_text(label)
        )

        return p

    for kind, value, url in parse_inline_segments(
        text
    ):

        if kind == 'text':

            add_mixed_text_runs(
                p,
                value
            )

        elif kind == 'link':

            # IMPORTANT: Only English text gets the hyperlink
            # Check if the entire link text is English only
            if is_english_only_text(value):
                # English text gets hyperlink
                color = english_color(value)

                add_external_hyperlink(
                    p,
                    value,
                    url,
                    color=color,
                    underline=False,
                    bold=True
                )
            else:
                # Split the link text into English and Arabic parts
                # English text gets hyperlink, Arabic text remains normal
                chunks = re.split(
                    r'([\u0600-\u06FF'
                    r'\u0750-\u077F'
                    r'\u08A0-\u08FF'
                    r'\uFB50-\uFDFF'
                    r'\uFE70-\uFEFF]+)',
                    value
                )

                for chunk in chunks:

                    if not chunk:
                        continue

                    # Check if this chunk is English only
                    if (
                        LATIN_RE.search(chunk)
                        and not ARABIC_RE.search(chunk)
                    ):
                        # English text gets hyperlink
                        color = english_color(chunk)

                        add_external_hyperlink(
                            p,
                            chunk,
                            url,
                            color=color,
                            underline=False,
                            bold=True
                        )
                    else:
                        # Arabic text remains normal
                        add_mixed_text_runs(
                            p,
                            chunk
                        )

        elif kind == 'code':

            r = p.add_run(value)

            set_run_font(
                r,
                'Consolas',
                size=9,
                color='C0392B'
            )

        elif kind == 'bold':

            for (
                subkind,
                subvalue,
                _
            ) in parse_inline_segments(value):

                if subkind == 'text':

                    chunks = re.split(
                        r'([\u0600-\u06FF'
                        r'\u0750-\u077F'
                        r'\u08A0-\u08FF'
                        r'\uFB50-\uFDFF'
                        r'\uFE70-\uFEFF]+)',
                        subvalue
                    )

                    for chunk in chunks:

                        if not chunk:
                            continue

                        color = (
                            english_color(chunk)
                            if (
                                LATIN_RE.search(chunk)
                                and
                                not ARABIC_RE.search(chunk)
                            )
                            else BODY_COLOR
                        )

                        r = p.add_run(
                            chunk
                        )

                        set_run_font(
                            r,
                            color=color,
                            bold=True,
                            size=11
                        )

                else:

                    r = p.add_run(
                        subvalue
                    )

                    set_run_font(
                        r,
                        bold=True,
                        size=11,
                        color=BODY_COLOR
                    )

        elif kind == 'italic':

            chunks = re.split(
                r'([\u0600-\u06FF'
                r'\u0750-\u077F'
                r'\u08A0-\u08FF'
                r'\uFB50-\uFDFF'
                r'\uFE70-\uFEFF]+)',
                value
            )

            for chunk in chunks:

                if not chunk:
                    continue

                color = (
                    english_color(chunk)
                    if (
                        LATIN_RE.search(chunk)
                        and
                        not ARABIC_RE.search(chunk)
                    )
                    else BODY_COLOR
                )

                r = p.add_run(
                    chunk
                )

                set_run_font(
                    r,
                    color=color,
                    italic=True,
                    size=11
                )

    return p


# ============================================================
# MULTIPLE "= URL" HYPERLINKS
# ============================================================

def add_multiple_url_hyperlinks_paragraph(
    doc,
    text
):
    """
    Converts multiple:

        Label = URL Label = URL Label = URL

    into:

        Label    Label    Label

    Each Label becomes a separate clickable hyperlink.

    IMPORTANT: Only English text gets the hyperlink. Arabic text remains as normal text.

    Example:

        Romanian Deadlift = URL1 النص العربي = URL2

    becomes:

        Romanian Deadlift النص العربي

    where:
        - Romanian Deadlift -> URL1 (hyperlink)
        - النص العربي -> normal text (no hyperlink)

    The '=' signs and URLs are NOT displayed.

    IMPORTANT:
    '=' characters inside URLs such as:

        ?si=XXXXXXXX

    are completely safe because we only detect:

        = https://
    """

    url_pattern = re.compile(
        r'\s*=\s*(https?://[^\s]+)',
        re.IGNORECASE
    )

    url_pattern_no_space = re.compile(
        r'\s*=(https?://[^\s]+)',
        re.IGNORECASE
    )

    matches = list(
        url_pattern.finditer(text)
    )

    # If no matches with space, try without space
    if not matches:
        matches = list(
            url_pattern_no_space.finditer(text)
        )

    if not matches:
        return False

    p = doc.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    p.paragraph_format.space_after = Pt(7)

    previous_url_end = 0

    for index, match in enumerate(matches):

        url = normalize_url(
            match.group(1)
        )

        # ----------------------------------------------------
        # Text between previous URL and current "= URL"
        # ----------------------------------------------------
        label = text[
            previous_url_end:
            match.start()
        ].strip()

        if label:

            # Check if the entire label is English only
            if is_english_only_text(label):
                # English text gets hyperlink
                if index > 0:
                    # Add space if not first element
                    spacer = p.add_run(' ')
                    set_run_font(
                        spacer,
                        size=11,
                        color=BODY_COLOR
                    )

                add_external_hyperlink(
                    p,
                    label,
                    url,
                    color=english_color(label),
                    underline=False,
                    bold=True
                )
            else:
                # Split label into English and Arabic parts
                # English text gets hyperlink, Arabic text remains normal
                chunks = re.split(
                    r'([\u0600-\u06FF'
                    r'\u0750-\u077F'
                    r'\u08A0-\u08FF'
                    r'\uFB50-\uFDFF'
                    r'\uFE70-\uFEFF]+)',
                    label
                )

                first_chunk = True

                for chunk in chunks:

                    if not chunk:
                        continue

                    # Check if this chunk is English only
                    if (
                        LATIN_RE.search(chunk)
                        and not ARABIC_RE.search(chunk)
                    ):
                        # English text gets hyperlink
                        if not first_chunk:
                            # Add space if not first element
                            spacer = p.add_run(' ')
                            set_run_font(
                                spacer,
                                size=11,
                                color=BODY_COLOR
                            )

                        add_external_hyperlink(
                            p,
                            chunk,
                            url,
                            color=english_color(chunk),
                            underline=False,
                            bold=True
                        )
                        first_chunk = False
                    else:
                        # Arabic text remains normal
                        if not first_chunk:
                            # Add space if not first element
                            spacer = p.add_run(' ')
                            set_run_font(
                                spacer,
                                size=11,
                                color=BODY_COLOR
                            )
                        add_mixed_text_runs(
                            p,
                            chunk
                        )
                        first_chunk = False

        previous_url_end = match.end()

    # --------------------------------------------------------
    # Any text after the final URL.
    # --------------------------------------------------------

    remaining = text[
        previous_url_end:
    ].strip()

    if remaining:

        spacer = p.add_run(' ')

        set_run_font(
            spacer,
            size=11,
            color=BODY_COLOR
        )

        add_mixed_text_runs(
            p,
            remaining
        )

    return True


# ============================================================
# Heading Numbering
# ============================================================

def make_heading_number(
    counters,
    level
):

    idx = level - 2

    counters[idx] += 1

    for j in range(
        idx + 1,
        len(counters)
    ):

        counters[j] = 0

    return '.'.join(
        ['0']
        +
        [
            f'{counters[j]:02d}'
            for j in range(idx + 1)
        ]
    )


def slugify(text):

    text = (
        text or ''
    ).strip().lower()

    text = re.sub(
        r'[^\w\s-]',
        '',
        text,
        flags=re.UNICODE
    )

    return re.sub(
        r'[\s_]+',
        '-',
        text
    ).strip('-')


def extract_headings(
    markdown_content
):

    headings = []

    counters = [
        0,
        0,
        0,
        0
    ]

    used = {}

    for line in markdown_content.splitlines():

        m = re.match(
            r'^\s*(#{2,5})\s+(.+?)\s*$',
            line
        )

        if not m:
            continue

        level = len(
            m.group(1)
        )

        title = m.group(2).strip()

        number = make_heading_number(
            counters,
            level
        )

        base = (
            f'heading-'
            f'{number.replace(".", "-")}-'
            f'{slugify(title) or "section"}'
        )

        count = (
            used.get(base, 0)
            + 1
        )

        used[base] = count

        bookmark = (
            base
            if count == 1
            else f'{base}-{count}'
        )

        headings.append(
            {
                'level': level,
                'number': number,
                'title': title,
                'bookmark': bookmark
            }
        )

    return headings


# ============================================================
# Tables
# ============================================================

def parse_table_block(lines):

    rows = []

    for line in lines:

        cells = [
            c.strip()
            for c in line.strip()
            .strip('|')
            .split('|')
        ]

        rows.append(cells)

    return rows


def add_table_mixed_text_runs(
    paragraph,
    text: str,
    *,
    bold=False,
    size=9,
    default_color=BODY_COLOR
):

    text = str(text)

    chunks = re.split(
        r'([\u0600-\u06FF'
        r'\u0750-\u077F'
        r'\u08A0-\u08FF'
        r'\uFB50-\uFDFF'
        r'\uFE70-\uFEFF]+)',
        text
    )

    for chunk in chunks:

        if not chunk:
            continue

        is_english = bool(
            LATIN_RE.search(chunk)
            and not ARABIC_RE.search(chunk)
        )

        base_color = (
            english_color(chunk)
            if is_english
            else default_color
        )

        pos = 0

        for match in QUESTION_MARKER_RE.finditer(
            chunk
        ):

            if match.start() > pos:

                before = chunk[
                    pos:match.start()
                ]

                run = paragraph.add_run(
                    before
                )

                set_run_font(
                    run,
                    size=size,
                    bold=(
                        bold
                        if not is_english
                        else False
                    ),
                    color=base_color
                )

            run = paragraph.add_run(
                match.group(0)
            )

            set_run_font(
                run,
                size=size,
                bold=(
                    bold
                    if not is_english
                    else False
                ),
                color=BLUE_ENGLISH
            )

            pos = match.end()

        if pos < len(chunk):

            run = paragraph.add_run(
                chunk[pos:]
            )

            set_run_font(
                run,
                size=size,
                bold=(
                    bold
                    if not is_english
                    else False
                ),
                color=base_color
            )


def add_table(
    doc,
    rows,
    table_index
):

    if not rows:
        return

    body_rows = rows

    if (
        len(rows) >= 2
        and all(
            set(cell.strip())
            <= set('-: ')
            for cell in rows[1]
        )
    ):

        body_rows = [
            rows[0]
        ] + rows[2:]

    cols = max(
        len(r)
        for r in body_rows
    )

    table = doc.add_table(
        rows=len(body_rows),
        cols=cols
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = 'Table Grid'

    table.autofit = True

    caption = doc.add_paragraph()

    caption.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = caption.add_run(
        f'Table {table_index}'
    )

    set_run_font(
        r,
        size=10,
        bold=True,
        color='1A5276'
    )

    for ri, row in enumerate(
        body_rows
    ):

        for ci in range(cols):

            cell = table.cell(
                ri,
                ci
            )

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            set_cell_border(
                cell
            )

            if ri == 0:

                set_cell_shading(
                    cell,
                    TABLE_HEADER_COLOR
                )

            else:

                if ri % 2 == 0:

                    set_cell_shading(
                        cell,
                        'EAF2F8'
                    )

            p = cell.paragraphs[0]

            p.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
            )

            p.clear()

            value = (
                row[ci]
                if ci < len(row)
                else ''
            )

            # First, check if the cell contains "= URL" pattern
            url_pattern = re.compile(
                r'\s*=\s*(https?://[^\s]+)',
                re.IGNORECASE
            )

            url_pattern_no_space = re.compile(
                r'\s*=(https?://[^\s]+)',
                re.IGNORECASE
            )

            url_matches = list(
                url_pattern.finditer(value)
            )

            # If no matches with space, try without space
            if not url_matches:
                url_matches = list(
                    url_pattern_no_space.finditer(value)
                )

            if url_matches:
                # Handle "= URL" pattern in table cells
                previous_url_end = 0

                for match in url_matches:

                    url = normalize_url(
                        match.group(1)
                    )

                    # Text before the "= URL"
                    label = value[
                        previous_url_end:
                        match.start()
                    ].strip()

                    if label:

                        # Check if the entire label is English only
                        if is_english_only_text(label):
                            # English text gets hyperlink
                            link_color = (
                                'FFFFFF'
                                if ri == 0
                                else (
                                    BLUE_ENGLISH
                                    if QUESTION_MARKER_RE.search(label)
                                    else english_color(label)
                                )
                            )

                            add_external_hyperlink(
                                p,
                                label,
                                url,
                                color=link_color,
                                underline=False,
                                bold=(ri == 0)
                            )
                        else:
                            # Split label into English and Arabic parts
                            chunks = re.split(
                                r'([\u0600-\u06FF'
                                r'\u0750-\u077F'
                                r'\u08A0-\u08FF'
                                r'\uFB50-\uFDFF'
                                r'\uFE70-\uFEFF]+)',
                                label
                            )

                            for chunk in chunks:

                                if not chunk:
                                    continue

                                if (
                                    LATIN_RE.search(chunk)
                                    and not ARABIC_RE.search(chunk)
                                ):
                                    # English text gets hyperlink
                                    link_color = (
                                        'FFFFFF'
                                        if ri == 0
                                        else (
                                            BLUE_ENGLISH
                                            if QUESTION_MARKER_RE.search(chunk)
                                            else english_color(chunk)
                                        )
                                    )

                                    add_external_hyperlink(
                                        p,
                                        chunk,
                                        url,
                                        color=link_color,
                                        underline=False,
                                        bold=(ri == 0)
                                    )
                                else:
                                    # Arabic text remains normal
                                    add_table_mixed_text_runs(
                                        p,
                                        chunk,
                                        bold=(ri == 0),
                                        size=9,
                                        default_color=(
                                            'FFFFFF'
                                            if ri == 0
                                            else BODY_COLOR
                                        )
                                    )

                    previous_url_end = match.end()

                # Any text after the final URL
                remaining = value[
                    previous_url_end:
                ].strip()

                if remaining:

                    add_table_mixed_text_runs(
                        p,
                        remaining,
                        bold=(ri == 0),
                        size=9,
                        default_color=(
                            'FFFFFF'
                            if ri == 0
                            else BODY_COLOR
                        )
                    )
            else:
                # Handle normal markdown links [text](url)
                for (
                    kind,
                    val,
                    url
                ) in parse_inline_segments(
                    value
                ):

                    if kind == 'text':

                        add_table_mixed_text_runs(
                            p,
                            val,
                            bold=(ri == 0),
                            size=9,
                            default_color=(
                                'FFFFFF'
                                if ri == 0
                                else BODY_COLOR
                            )
                        )

                    elif kind == 'link':

                        # IMPORTANT: Only English text gets the hyperlink
                        # Check if the entire link text is English only
                        if is_english_only_text(val):
                            # English text gets hyperlink
                            link_color = (
                                'FFFFFF'
                                if ri == 0
                                else (
                                    BLUE_ENGLISH
                                    if QUESTION_MARKER_RE.search(val)
                                    else english_color(val)
                                )
                            )

                            add_external_hyperlink(
                                p,
                                val,
                                url,
                                color=link_color,
                                underline=False,
                                bold=(ri == 0)
                            )
                        else:
                            # Split the link text into English and Arabic parts
                            # English text gets hyperlink, Arabic text remains normal
                            chunks = re.split(
                                r'([\u0600-\u06FF'
                                r'\u0750-\u077F'
                                r'\u08A0-\u08FF'
                                r'\uFB50-\uFDFF'
                                r'\uFE70-\uFEFF]+)',
                                val
                            )

                            for chunk in chunks:

                                if not chunk:
                                    continue

                                # Check if this chunk is English only
                                if (
                                    LATIN_RE.search(chunk)
                                    and not ARABIC_RE.search(chunk)
                                ):
                                    # English text gets hyperlink
                                    link_color = (
                                        'FFFFFF'
                                        if ri == 0
                                        else (
                                            BLUE_ENGLISH
                                            if QUESTION_MARKER_RE.search(chunk)
                                            else english_color(chunk)
                                        )
                                    )

                                    add_external_hyperlink(
                                        p,
                                        chunk,
                                        url,
                                        color=link_color,
                                        underline=False,
                                        bold=(ri == 0)
                                    )
                                else:
                                    # Arabic text remains normal
                                    add_table_mixed_text_runs(
                                        p,
                                        chunk,
                                        bold=(ri == 0),
                                        size=9,
                                        default_color=(
                                            'FFFFFF'
                                            if ri == 0
                                            else BODY_COLOR
                                        )
                                    )

                    elif kind == 'code':

                        rr = p.add_run(val)

                        set_run_font(
                            rr,
                            'Consolas',
                            size=8,
                            color=(
                                'FFFFFF'
                                if ri == 0
                                else BODY_COLOR
                            )
                        )

                    else:

                        add_table_mixed_text_runs(
                            p,
                            val,
                            bold=(ri == 0),
                            size=9,
                            default_color=(
                                'FFFFFF'
                                if ri == 0
                                else BODY_COLOR
                            )
                        )

    set_repeat_table_header(
        table.rows[0]
    )

    return table


# ============================================================
# Info Box / Quote / Lists / Code
# ============================================================

def add_info_box(
    doc,
    title,
    body_lines
):

    table = doc.add_table(
        rows=1,
        cols=1
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.autofit = True

    cell = table.cell(
        0,
        0
    )

    set_cell_shading(
        cell,
        'FEF9E7'
    )

    set_cell_border(
        cell,
        color='F9E79F',
        sz='8'
    )

    p = cell.paragraphs[0]

    p.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    r = p.add_run(title)

    set_run_font(
        r,
        size=11,
        bold=True,
        color='D68910'
    )

    for line in body_lines:

        p = cell.add_paragraph()

        p.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT
        )

        add_mixed_text_runs(
            p,
            line
        )

    doc.add_paragraph()


def add_quote(
    doc,
    text
):

    table = doc.add_table(
        rows=1,
        cols=1
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    cell = table.cell(
        0,
        0
    )

    set_cell_shading(
        cell,
        'EAF2F8'
    )

    set_cell_border(
        cell,
        color='2980B9',
        sz='10'
    )

    p = cell.paragraphs[0]

    p.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    add_mixed_text_runs(
        p,
        text
    )

    doc.add_paragraph()


def add_list_item(
    doc,
    text,
    ordered=False,
    level=0
):

    style = (
        'List Number'
        if ordered
        else 'List Bullet'
    )

    p = doc.add_paragraph(
        style=style
    )

    p.paragraph_format.left_indent = (
        Inches(0.25 * level)
    )

    p.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    # Check if the text contains "= URL" pattern
    url_pattern = re.compile(
        r'\s*=\s*(https?://[^\s]+)',
        re.IGNORECASE
    )

    url_pattern_no_space = re.compile(
        r'\s*=(https?://[^\s]+)',
        re.IGNORECASE
    )

    url_matches = list(
        url_pattern.finditer(text)
    )

    # If no matches with space, try without space
    if not url_matches:
        url_matches = list(
            url_pattern_no_space.finditer(text)
        )

    if url_matches:
        # Handle "= URL" pattern in list items
        previous_url_end = 0

        for match in url_matches:

            url = normalize_url(
                match.group(1)
            )

            # Text before the "= URL"
            label = text[
                previous_url_end:
                match.start()
            ].strip()

            if label:

                # Check if the entire label is English only
                if is_english_only_text(label):
                    # English text gets hyperlink
                    add_external_hyperlink(
                        p,
                        label,
                        url,
                        color=english_color(label),
                        underline=False,
                        bold=True
                    )
                else:
                    # Split label into English and Arabic parts
                    chunks = re.split(
                        r'([\u0600-\u06FF'
                        r'\u0750-\u077F'
                        r'\u08A0-\u08FF'
                        r'\uFB50-\uFDFF'
                        r'\uFE70-\uFEFF]+)',
                        label
                    )

                    for chunk in chunks:

                        if not chunk:
                            continue

                        if (
                            LATIN_RE.search(chunk)
                            and not ARABIC_RE.search(chunk)
                        ):
                            # English text gets hyperlink
                            add_external_hyperlink(
                                p,
                                chunk,
                                url,
                                color=english_color(chunk),
                                underline=False,
                                bold=True
                            )
                        else:
                            # Arabic text remains normal
                            add_mixed_text_runs(
                                p,
                                chunk
                            )

            previous_url_end = match.end()

        # Any text after the final URL
        remaining = text[
            previous_url_end:
        ].strip()

        if remaining:

            add_mixed_text_runs(
                p,
                remaining
            )
    else:
        # Handle normal text without "= URL" pattern
        add_mixed_text_runs(
            p,
            text
        )

    return p


def add_code_block(
    doc,
    code,
    language='text'
):

    p = doc.add_paragraph(
        style='Code Block'
    )

    p.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    p.paragraph_format.space_after = Pt(7)

    r = p.add_run(code)

    set_run_font(
        r,
        'Consolas',
        size=8.5,
        color=BODY_COLOR
    )

    pPr = p._p.get_or_add_pPr()

    shd = OxmlElement(
        'w:shd'
    )

    shd.set(
        qn('w:fill'),
        'F8F9FA'
    )

    pPr.append(shd)

    return p


# ============================================================
# Mermaid
# ============================================================

def render_mermaid_to_png(
    raw_code,
    output_png
):

    """
    Best-effort Mermaid -> PNG using
    Playwright + Mermaid CDN.
    """

    try:

        from playwright.sync_api import (
            sync_playwright
        )

        import cairosvg

        html_doc = f'''
        <!doctype html>
        <html>
        <head>
            <meta charset="utf-8">
        </head>
        <body>
        <pre class="mermaid">
        {html.escape(raw_code)}
        </pre>

        <script src=
        "https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.min.js">
        </script>

        </body>
        </html>
        '''

        with sync_playwright() as p:

            browser = p.chromium.launch()

            page = browser.new_page(
                viewport={
                    'width': 1400,
                    'height': 1000
                }
            )

            page.set_content(
                html_doc,
                wait_until='networkidle'
            )

            page.wait_for_function(
                "() => typeof mermaid !== 'undefined'"
            )

            page.evaluate(
                '''async () => {
                    mermaid.initialize({
                        startOnLoad: false,
                        securityLevel: 'loose',
                        theme: 'default'
                    });

                    const node =
                        document.querySelector('.mermaid');

                    const result =
                        await mermaid.render(
                            'word-mermaid',
                            node.textContent
                        );

                    node.innerHTML =
                        result.svg;
                }'''
            )

            svg = page.locator(
                'svg'
            ).evaluate(
                '(el) => el.outerHTML'
            )

            browser.close()

        cairosvg.svg2png(
            bytestring=svg.encode('utf-8'),
            write_to=str(output_png),
            output_width=1400
        )

        return True

    except Exception as exc:

        print(
            '  WARNING: Mermaid could not be '
            f'rendered into Word image: {exc}'
        )

        return False


# ============================================================
# Document Title
# ============================================================

def add_document_title(
    doc,
    title
):

    p = doc.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(title)

    set_run_font(
        r,
        size=22,
        bold=True,
        color='1A5276'
    )

    p.paragraph_format.space_after = Pt(18)

    return p


# ============================================================
# Extract Title
# ============================================================

def extract_title(
    markdown_content,
    fallback
):

    for line in markdown_content.splitlines():

        m = re.match(
            r'^\s*#\s+(.+?)\s*$',
            line
        )

        if m:

            return m.group(1).strip()

    return fallback


# ============================================================
# Build DOCX
# ============================================================

def build_document(
    markdown_content,
    source_name
):

    title = extract_title(
        markdown_content,
        Path(source_name).stem
    )

    headings = extract_headings(
        markdown_content
    )

    doc = Document()

    configure_styles(doc)

    set_update_fields_on_open(
        doc
    )

    section = doc.sections[0]

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    section.different_first_page_header_footer = False

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    add_document_title(
        doc,
        title
    )

    content_section = doc.sections[0]

    set_section_page_number_restart(
        content_section,
        1
    )

    add_footer_page_number(
        content_section
    )

    # --------------------------------------------------------
    # Rendering state
    # --------------------------------------------------------

    counters = [
        0,
        0,
        0,
        0
    ]

    table_index = 0

    lines = markdown_content.splitlines()

    i = 0

    n = len(lines)

    bookmark_id = 100

    current_heading_index = 0

    # ========================================================
    # Main Markdown Loop
    # ========================================================

    while i < n:

        raw = lines[i]

        stripped = raw.strip()

        if not stripped:

            i += 1

            continue

        # ----------------------------------------------------
        # H1 = document title
        # ----------------------------------------------------

        if (
            re.match(
                r'^#\s+',
                stripped
            )
            and not re.match(
                r'^##\s+',
                stripped
            )
        ):

            i += 1

            continue

        # ----------------------------------------------------
        # Heading H2-H5
        # ----------------------------------------------------

        hm = re.match(
            r'^(#{2,5})\s+(.+?)\s*$',
            stripped
        )

        if hm:

            level = len(
                hm.group(1)
            )

            text = hm.group(2).strip()

            number = make_heading_number(
                counters,
                level
            )

            if (
                current_heading_index
                <
                len(headings)
            ):

                h = headings[
                    current_heading_index
                ]

                current_heading_index += 1

            else:

                h = {
                    'level': level,
                    'number': number,
                    'title': text,
                    'bookmark': (
                        f'heading-{bookmark_id}'
                    )
                }

            style_name = (
                f'Heading {level - 1}'
            )

            p = doc.add_paragraph(
                style=style_name
            )

            p.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
            )

            p.paragraph_format.keep_with_next = True

            p.paragraph_format.space_before = Pt(
                12
                if level == 2
                else 8
            )

            p.paragraph_format.space_after = Pt(5)

            add_bookmark(
                p,
                h['bookmark'],
                bookmark_id
            )

            bookmark_id += 1

            rn = p.add_run(
                h['number'] + ' '
            )

            set_run_font(
                rn,
                size={
                    2: 15,
                    3: 13,
                    4: 11,
                    5: 10
                }[level],
                bold=True,
                color=HEADING_COLORS[
                    level - 1
                ]
            )

            add_mixed_text_runs(
                p,
                text
            )

            i += 1

            continue

        # ----------------------------------------------------
        # Code block
        # ----------------------------------------------------

        if stripped.startswith('```'):

            language = (
                stripped[3:].strip()
                or 'text'
            )

            code_lines = []

            i += 1

            while (
                i < n
                and not lines[i]
                .strip()
                .startswith('```')
            ):

                code_lines.append(
                    lines[i]
                )

                i += 1

            raw_code = '\n'.join(
                code_lines
            ).strip('\n')

            if language.lower() == 'mermaid':

                with tempfile.TemporaryDirectory() as td:

                    png = (
                        Path(td)
                        /
                        'mermaid.png'
                    )

                    if render_mermaid_to_png(
                        raw_code,
                        png
                    ):

                        p = doc.add_paragraph()

                        p.alignment = (
                            WD_ALIGN_PARAGRAPH.CENTER
                        )

                        r = p.add_run()

                        r.add_picture(
                            str(png),
                            width=Inches(6.4)
                        )

                    else:

                        add_code_block(
                            doc,
                            raw_code,
                            language
                        )

            else:

                add_code_block(
                    doc,
                    raw_code,
                    language
                )

            i += 1

            continue

        # ----------------------------------------------------
        # Markdown table
        # ----------------------------------------------------

        if (
            '|'
            in stripped
            and i + 1 < n
            and '|'
            in lines[i + 1]
            and set(
                lines[i + 1].strip()
            )
            -
            set('|-: ')
            == set()
        ):

            table_lines = []

            while (
                i < n
                and '|'
                in lines[i]
            ):

                table_lines.append(
                    lines[i]
                )

                i += 1

            table_index += 1

            add_table(
                doc,
                parse_table_block(
                    table_lines
                ),
                table_index
            )

            continue

        # ----------------------------------------------------
        # Blockquote
        # ----------------------------------------------------

        if stripped.startswith('> '):

            quote_lines = []

            while (
                i < n
                and lines[i]
                .strip()
                .startswith('> ')
            ):

                quote_lines.append(
                    lines[i]
                    .strip()[2:]
                )

                i += 1

            add_quote(
                doc,
                ' '.join(quote_lines)
            )

            continue

        # ----------------------------------------------------
        # Lists
        # ----------------------------------------------------

        if (
            stripped.startswith((
                '- ',
                '* '
            ))
            or re.match(
                r'^\d+\. ',
                stripped
            )
        ):

            ordered = bool(
                re.match(
                    r'^\d+\. ',
                    stripped
                )
            )

            while i < n:

                candidate = lines[i]

                cs = candidate.strip()

                if not cs:

                    break

                if not ordered:

                    m = re.match(
                        r'^[-*]\s+(.+)$',
                        cs
                    )

                else:

                    m = re.match(
                        r'^\d+\.\s+(.+)$',
                        cs
                    )

                if m:

                    add_list_item(
                        doc,
                        m.group(1),
                        ordered=ordered
                    )

                    i += 1

                elif (
                    candidate.startswith('  ')
                    or candidate.startswith('\t')
                ):

                    add_list_item(
                        doc,
                        cs,
                        ordered=ordered,
                        level=1
                    )

                    i += 1

                else:

                    break

            continue

        # ----------------------------------------------------
        # Info boxes
        # ----------------------------------------------------

        if stripped in [
            'Important Notes',
            'Warnings',
            'Tips',
            'Recommendations'
        ]:

            box_title = stripped

            body = []

            i += 1

            while (
                i < n
                and lines[i].strip()
            ):

                body.append(
                    lines[i].strip()
                )

                i += 1

            add_info_box(
                doc,
                box_title,
                body
            )

            continue

        # ====================================================
        # HYPERLINK SYSTEM
        # ====================================================

        # ----------------------------------------------------
        # Case 1:
        #
        # Multiple hyperlinks in ONE line:
        #
        # Romanian Deadlift = URL1 Romanian Deadlift = URL2
        #
        # Each sentence receives its own URL.
        # ----------------------------------------------------

        if re.search(
            r'=\s*https?://[^\s]+',
            stripped,
            re.IGNORECASE
        ) or re.search(
            r'=https?://[^\s]+',
            stripped,
            re.IGNORECASE
        ):

            if add_multiple_url_hyperlinks_paragraph(
                doc,
                stripped
            ):

                i += 1

                continue

        # ----------------------------------------------------
        # Case 2:
        #
        # Sentence
        # = https://example.com
        #
        # Attach URL to previous paragraph.
        # IMPORTANT: Only English text gets the hyperlink.
        # ----------------------------------------------------

        standalone = re.match(
            r'^=\s*(https?://[^\s]+)\s*$',
            stripped,
            re.IGNORECASE
        )

        if standalone:

            url = normalize_url(
                standalone.group(1)
            )

            if doc.paragraphs:

                previous = doc.paragraphs[-1]

                visible = ''.join(
                    r.text or ''
                    for r in previous.runs
                ).strip()

                if visible:

                    for r in list(
                        previous.runs
                    ):

                        r._element.getparent().remove(
                            r._element
                        )

                    # Check if the entire visible text is English only
                    if is_english_only_text(visible):
                        # English text gets hyperlink
                        add_external_hyperlink(
                            previous,
                            visible,
                            url,
                            color=english_color(visible),
                            underline=False,
                            bold=True
                        )
                    else:
                        # Split visible text into English and Arabic parts
                        # English text gets hyperlink, Arabic text remains normal
                        chunks = re.split(
                            r'([\u0600-\u06FF'
                            r'\u0750-\u077F'
                            r'\u08A0-\u08FF'
                            r'\uFB50-\uFDFF'
                            r'\uFE70-\uFEFF]+)',
                            visible
                        )

                        for chunk in chunks:

                            if not chunk:
                                continue

                            # Check if this chunk is English only
                            if (
                                LATIN_RE.search(chunk)
                                and not ARABIC_RE.search(chunk)
                            ):
                                # English text gets hyperlink
                                add_external_hyperlink(
                                    previous,
                                    chunk,
                                    url,
                                    color=english_color(chunk),
                                    underline=False,
                                    bold=True
                                )
                            else:
                                # Arabic text remains normal
                                add_mixed_text_runs(
                                    previous,
                                    chunk
                                )

            i += 1

            continue

        # ----------------------------------------------------
        # Case 2b:
        #
        # Sentence
        # =https://example.com (no space after =)
        #
        # Attach URL to previous paragraph.
        # IMPORTANT: Only English text gets the hyperlink.
        # ----------------------------------------------------

        standalone_no_space = re.match(
            r'^=(https?://[^\s]+)\s*$',
            stripped,
            re.IGNORECASE
        )

        if standalone_no_space:

            url = normalize_url(
                standalone_no_space.group(1)
            )

            if doc.paragraphs:

                previous = doc.paragraphs[-1]

                visible = ''.join(
                    r.text or ''
                    for r in previous.runs
                ).strip()

                if visible:

                    for r in list(
                        previous.runs
                    ):

                        r._element.getparent().remove(
                            r._element
                        )

                    # Check if the entire visible text is English only
                    if is_english_only_text(visible):
                        # English text gets hyperlink
                        add_external_hyperlink(
                            previous,
                            visible,
                            url,
                            color=english_color(visible),
                            underline=False,
                            bold=True
                        )
                    else:
                        # Split visible text into English and Arabic parts
                        # English text gets hyperlink, Arabic text remains normal
                        chunks = re.split(
                            r'([\u0600-\u06FF'
                            r'\u0750-\u077F'
                            r'\u08A0-\u08FF'
                            r'\uFB50-\uFDFF'
                            r'\uFE70-\uFEFF]+)',
                            visible
                        )

                        for chunk in chunks:

                            if not chunk:
                                continue

                            # Check if this chunk is English only
                            if (
                                LATIN_RE.search(chunk)
                                and not ARABIC_RE.search(chunk)
                            ):
                                # English text gets hyperlink
                                add_external_hyperlink(
                                    previous,
                                    chunk,
                                    url,
                                    color=english_color(chunk),
                                    underline=False,
                                    bold=True
                                )
                            else:
                                # Arabic text remains normal
                                add_mixed_text_runs(
                                    previous,
                                    chunk
                                )

            i += 1

            continue

        # ----------------------------------------------------
        # Case 3:
        #
        # Normal Markdown paragraph.
        # ----------------------------------------------------

        add_rich_paragraph(
            doc,
            stripped
        )

        i += 1

    # ========================================================
    # Top / Bottom bookmarks
    # ========================================================

    if doc.paragraphs:

        add_bookmark(
            doc.paragraphs[0],
            'ContentTop',
            9000
        )

        add_bookmark(
            doc.paragraphs[-1],
            'ContentBottom',
            9001
        )

    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    for sec in doc.sections[1:]:

        add_footer_page_number(
            sec
        )

    return doc, headings


# ============================================================
# Safe DOCX Save
# ============================================================

def save_docx_safely(
    doc,
    output_docx
):

    output_docx = Path(
        output_docx
    )

    output_docx.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    temp_docx = output_docx.with_name(
        f'.{output_docx.stem}'
        f'_temp_{os.getpid()}.docx'
    )

    try:

        doc.save(
            temp_docx
        )

        if output_docx.exists():

            try:

                output_docx.unlink()

            except PermissionError:

                raise PermissionError(
                    'Cannot replace existing file '
                    'because it is locked:\n'
                    f'{output_docx}\n\n'
                    'Close the DOCX in Microsoft '
                    'Word/OneDrive and run again.'
                )

        temp_docx.replace(
            output_docx
        )

        return True

    except PermissionError:

        if temp_docx.exists():

            try:
                temp_docx.unlink()
            except Exception:
                pass

        raise

    except Exception:

        if temp_docx.exists():

            try:
                temp_docx.unlink()
            except Exception:
                pass

        raise


# ============================================================
# DOCX -> PDF
# ============================================================

def convert_docx_to_pdf_with_navigation(
    docx_path,
    pdf_path,
    headings
):

    word = None

    document = None

    try:

        docx_path = Path(
            docx_path
        ).resolve()

        pdf_path = Path(
            pdf_path
        ).resolve()

        print(
            '  Converting DOCX to PDF...'
        )

        # ----------------------------------------------------
        # STEP 1
        # DOCX -> PDF using Microsoft Word
        # ----------------------------------------------------

        word = win32com.client.DispatchEx(
            'Word.Application'
        )

        word.Visible = False

        word.DisplayAlerts = False

        document = word.Documents.Open(
            str(docx_path),
            ReadOnly=True
        )

        try:

            document.Fields.Update()

        except Exception:

            pass

        document.ExportAsFixedFormat(
            OutputFileName=str(pdf_path),

            ExportFormat=17,

            OpenAfterExport=False,

            OptimizeFor=0,

            Range=0,

            From=0,

            To=0,

            Item=0,

            IncludeDocProps=True,

            KeepIRM=True,

            CreateBookmarks=1,

            DocStructureTags=True,

            BitmapMissingFonts=True,

            UseISO19005_1=False
        )

        document.Close(False)

        document = None

        word.Quit()

        word = None

        print(
            f'  PDF created: {pdf_path}'
        )

        # ----------------------------------------------------
        # STEP 2
        # Build PDF bookmarks using PyMuPDF
        # ----------------------------------------------------

        try:

            add_pdf_navigation_bookmarks(
                pdf_path,
                headings
            )

            print(
                '  PDF Navigation/Bookmarks: ENABLED'
            )

        except Exception as ex:

            print(
                '  WARNING: PDF was created, '
                'but custom navigation could '
                f'not be added: {ex}'
            )

        return pdf_path

    except Exception as ex:

        print(
            '  ERROR converting DOCX to PDF: '
            f'{ex}'
        )

        return None

    finally:

        if document is not None:

            try:

                document.Close(False)

            except Exception:

                pass

        if word is not None:

            try:

                word.Quit()

            except Exception:

                pass


# ============================================================
# PDF Navigation / Bookmarks
# ============================================================

def add_pdf_navigation_bookmarks(
    pdf_path,
    headings
):

    """
    Create PDF Outline / Bookmarks.

    Markdown:

        H2 -> PDF level 1
        H3 -> PDF level 2
        H4 -> PDF level 3
        H5 -> PDF level 4

    The hierarchy is normalized so that PyMuPDF
    always receives a valid outline.
    """

    pdf_path = Path(
        pdf_path
    )

    pdf = fitz.open(
        str(pdf_path)
    )

    try:

        if not headings:

            print(
                '  No headings found. '
                'PDF navigation was not created.'
            )

            return

        toc = []

        last_search_page = 0

        # ----------------------------------------------------
        # Build preliminary TOC
        # ----------------------------------------------------

        for h in headings:

            title = (
                f"{h['number']} "
                f"{h['title']}"
            )

            matches = []

            # ------------------------------------------------
            # Exact numbered heading search
            # ------------------------------------------------

            for page_index in range(
                last_search_page,
                len(pdf)
            ):

                page = pdf[
                    page_index
                ]

                found = page.search_for(
                    title
                )

                if found:

                    matches = found

                    last_search_page = (
                        page_index
                    )

                    break

            # ------------------------------------------------
            # Fallback: title only
            # ------------------------------------------------

            if not matches:

                for page_index in range(
                    last_search_page,
                    len(pdf)
                ):

                    page = pdf[
                        page_index
                    ]

                    found = page.search_for(
                        h['title']
                    )

                    if found:

                        matches = found

                        last_search_page = (
                            page_index
                        )

                        break

            # ------------------------------------------------
            # If not found
            # ------------------------------------------------

            if matches:

                page_number = (
                    last_search_page + 1
                )

            else:

                page_number = 1

                print(
                    '  WARNING: Heading not '
                    'located in PDF: '
                    f'{title}'
                )

            requested_level = (
                h['level'] - 1
            )

            requested_level = max(
                1,
                min(
                    requested_level,
                    4
                )
            )

            toc.append(
                [
                    requested_level,
                    title,
                    page_number
                ]
            )

        # ----------------------------------------------------
        # Normalize hierarchy
        #
        # Invalid:
        #
        # 1
        # 3
        #
        # Valid:
        #
        # 1
        # 2
        # ----------------------------------------------------

        normalized_toc = []

        previous_level = 0

        for (
            level,
            title,
            page_number
        ) in toc:

            if previous_level == 0:

                level = 1

            elif level > (
                previous_level + 1
            ):

                level = (
                    previous_level + 1
                )

            level = max(
                1,
                level
            )

            normalized_toc.append(
                [
                    level,
                    title,
                    page_number
                ]
            )

            previous_level = level

        # ----------------------------------------------------
        # GUARANTEE:
        # first item must be level 1
        # ----------------------------------------------------

        if normalized_toc:

            normalized_toc[0][0] = 1

        # ----------------------------------------------------
        # Set outline
        # ----------------------------------------------------

        pdf.set_toc(
            normalized_toc
        )

        # ----------------------------------------------------
        # Safe PDF save
        # ----------------------------------------------------

        temp_pdf = pdf_path.with_name(
            f'.{pdf_path.stem}'
            f'_navigation_temp.pdf'
        )

        pdf.save(
            str(temp_pdf),
            garbage=4,
            deflate=True
        )

        pdf.close()

        if pdf_path.exists():

            pdf_path.unlink()

        temp_pdf.replace(
            pdf_path
        )

    except Exception:

        try:

            pdf.close()

        except Exception:

            pass

        raise


# ============================================================
# Process One File
# ============================================================

def process_file(
    input_file
):

    input_file = Path(
        input_file
    )

    if input_file.suffix.lower() != '.md':

        raise ValueError(
            f'Not a Markdown file: '
            f'{input_file}'
        )

    output_dir = (
        input_file.parent
        /
        'WORD'
    )

    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    print(
        f'\nConverting to Word: '
        f'{input_file}'
    )

    md_content = read_markdown_file(
        input_file
    )

    doc, headings = build_document(
        md_content,
        input_file.name
    )

    # --------------------------------------------------------
    # Save DOCX
    # --------------------------------------------------------

    output_docx = (
        output_dir
        /
        f'{input_file.stem}.docx'
    )

    try:

        save_docx_safely(
            doc,
            output_docx
        )

    except PermissionError as ex:

        print(
            f'  ERROR: {ex}'
        )

        return None

    print(
        f'  Sections: {len(headings)}'
    )

    print(
        f'  DOCX saved to: '
        f'{output_docx}'
    )

    # --------------------------------------------------------
    # Convert DOCX -> PDF
    # --------------------------------------------------------

    output_pdf = (
        output_dir
        /
        f'{input_file.stem}.pdf'
    )

    pdf_result = (
        convert_docx_to_pdf_with_navigation(
            output_docx,
            output_pdf,
            headings
        )
    )

    if pdf_result:

        print(
            f'  PDF saved to: '
            f'{output_pdf}'
        )

        print(
            '  PDF Navigation/Bookmarks: ENABLED'
        )

    return output_docx


# ============================================================
# Main
# ============================================================

def main(
    input_path
):

    input_path = Path(
        input_path
    )

    if not input_path.exists():

        print(
            f'ERROR: Path does not exist: '
            f'{input_path}'
        )

        sys.exit(1)

    # --------------------------------------------------------
    # Single file
    # --------------------------------------------------------

    if input_path.is_file():

        if (
            input_path.suffix.lower()
            != '.md'
        ):

            print(
                'ERROR: Input file must '
                'be a .md file.'
            )

            sys.exit(1)

        process_file(
            input_path
        )

        return

    # --------------------------------------------------------
    # Folder
    # --------------------------------------------------------

    if input_path.is_dir():

        md_files = sorted(
            p
            for p in input_path.rglob(
                '*.md'
            )
            if p.is_file()
        )

        if not md_files:

            print(
                f'No .md files found under: '
                f'{input_path}'
            )

            return

        print(
            f'Found {len(md_files)} '
            'Markdown files.'
        )
        for md_file in md_files:
            try:
                process_file(md_file)
            except Exception as ex:
                print(f'  ERROR processing {md_file}: {ex}')
        print('\nDone.')
        return

    print('ERROR: Input path must be a file or folder.')
    sys.exit(1)


if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(
            'Usage:\n'
            '  py generatePdfTablesAndIndexingFromMd.py <markdown-file-or-folder>\n\n'
            'Examples:\n'
            r'  py generatePdfTablesAndIndexingFromMd.py "E:\\mazen\\English\\Authentication.md"' '\n'
            r'  py generatePdfTablesAndIndexingFromMd.py "E:\\mazen\\English"'
        )
        sys.exit(1)
    main(sys.argv[1])